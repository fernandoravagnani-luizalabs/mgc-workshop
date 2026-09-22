#!/usr/bin/env python3
"""
Gera a apresentação em PowerPoint (.pptx) para o Workshop Magalu Cloud na Prática.
Design limpo, fundo branco, widescreen 16:9, cards e caixas de código legíveis.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Dimensões 16:9 widescreen (13.333 x 7.5 polegadas)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Paleta de cores sóbria e corporativa (fundo branco)
COLOR_BG_WHITE = RGBColor(255, 255, 255)
COLOR_CARD_BG = RGBColor(248, 250, 252)       # Slate 50 suave
COLOR_CARD_BORDER = RGBColor(226, 232, 240)   # Slate 200
COLOR_MAGALU_BLUE = RGBColor(0, 134, 255)     # Azul Magalu #0086FF
COLOR_PRIMARY_DARK = RGBColor(15, 23, 42)     # Slate 900 (Títulos)
COLOR_TEXT_BODY = RGBColor(51, 65, 85)        # Slate 700 (Corpo)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)    # Slate 500 (Secundário/Legenda)
COLOR_CODE_BG = RGBColor(241, 245, 249)       # Fundo para código
COLOR_CODE_TEXT = RGBColor(15, 23, 42)        # Texto código
COLOR_ACCENT_GREEN = RGBColor(16, 185, 129)   # Verde Sucesso
COLOR_ACCENT_ORANGE = RGBColor(245, 158, 11)  # Laranja Alerta

FONT_HEADING = "Trebuchet MS"
FONT_BODY = "Calibri"
FONT_CODE = "Consolas"

def create_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank_slide_layout = prs.slide_layouts[6] # Blank layout

    def add_blank_slide():
        slide = prs.slides.add_slide(blank_slide_layout)
        # Garante fundo branco explícito
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG_WHITE
        return slide

    def add_header(slide, title_text, category="WORKSHOP MAGALU CLOUD"):
        # Categoria / Tag superior
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.3))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.name = FONT_HEADING
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = COLOR_MAGALU_BLUE

        # Título principal
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.6))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY_DARK

    def add_footer(slide, current_slide, total_slides=22):
        # Linha / texto discreto no rodapé
        f_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.3))
        tf = f_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Magalu Cloud na Prática  |  docs.magalu.cloud  |  Slide {current_slide} de {total_slides}"
        p.font.name = FONT_BODY
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)

        # Adicionar texto
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY_DARK
        p0.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.name = FONT_BODY
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_BODY
            p.space_after = Pt(6)

    def add_code_block(slide, left, top, width, height, title, code_lines):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CODE_BG
        shape.line.color.rgb = COLOR_CARD_BORDER
        shape.line.width = Pt(1)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_MAGALU_BLUE
        p0.space_after = Pt(8)

        for line in code_lines:
            p = tf.add_paragraph()
            p.text = line
            p.font.name = FONT_CODE
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_CODE_TEXT
            p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 1: Capa
    # =========================================================================
    s1 = add_blank_slide()
    # Barra lateral ou detalhe em azul Magalu
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_MAGALU_BLUE
    bar.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p_badge = tf1.paragraphs[0]
    p_badge.text = "WORKSHOP HANDS-ON (1H30)"
    p_badge.font.name = FONT_HEADING
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_MAGALU_BLUE
    p_badge.space_after = Pt(8)

    p_title = tf1.add_paragraph()
    p_title.text = "Magalu Cloud na Prática"
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(38)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY_DARK
    p_title.space_after = Pt(12)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Construindo e Automatizando Infraestrutura de Nuvem com Compute, Object Storage S3 e IaC"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = COLOR_TEXT_BODY
    p_sub.space_after = Pt(20)

    p_author = tf1.add_paragraph()
    p_author.text = "Magalu Cloud (MGC)  •  Documentação: docs.magalu.cloud"
    p_author.font.name = FONT_BODY
    p_author.font.size = Pt(13)
    p_author.font.color.rgb = COLOR_TEXT_MUTED

    add_footer(s1, 1)

    # =========================================================================
    # SLIDE 2: Objetivos do Workshop
    # =========================================================================
    s2 = add_blank_slide()
    add_header(s2, "Objetivos e Dinâmica do Workshop", "VISÃO GERAL")
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "🎯 Propósito do Workshop", [
        "Apresentar a plataforma de nuvem pública 100% brasileira da Magalu Cloud.",
        "Simular um caso real: hospedar um website corporativo escalável na nuvem.",
        "Experiência Hands-On individual: cada participante cria, testa e apaga seu ambiente.",
        "Modelo pay-as-you-go: provisionar e depois desmontar tudo sem custo residual."
    ])
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "⏱️ Agenda (90 Minutos)", [
        "15 min: Apresentação da Magalu Cloud e Arquitetura da Solução.",
        "15 min: Criação da Conta no ID Magalu e Pré-requisitos.",
        "25 min: Provisionamento da Infraestrutura (Console, CLI ou IaC).",
        "20 min: Deploy da Aplicação, Integração S3 e Validação de Mídias.",
        "15 min: Desmontagem Completa dos Recursos e Sessão de Perguntas (Q&A)."
    ])
    add_footer(s2, 2)

    # =========================================================================
    # SLIDE 3: Arquitetura da Solução
    # =========================================================================
    s3 = add_blank_slide()
    add_header(s3, "Arquitetura da Solução: WordPress + Object Storage", "ARQUITETURA")
    add_card(s3, Inches(0.8), Inches(1.6), Inches(3.6), Inches(4.8), "1. Compute & Rede", [
        "1 Máquina Virtual (VM) com Ubuntu 24.04 LTS.",
        "Flavor padrão BV1-1-10 (1 vCPU, 1GB RAM, 10GB Disco) ou BV2-2-8.",
        "1 IP Público IPv4 dedicado para acesso via navegador.",
        "Security Group com portas 22 (SSH) e 80 (HTTP) liberadas."
    ])
    add_card(s3, Inches(4.8), Inches(1.6), Inches(3.6), Inches(4.8), "2. Aplicação & Dados", [
        "WordPress oficial rodando em container Docker (Apache + PHP).",
        "Banco de dados relacional MySQL 8.0 em container dedicado.",
        "Deploy automatizado via Docker Compose na VM.",
        "Código, temas e configurações isolados no disco da VM."
    ])
    add_card(s3, Inches(8.8), Inches(1.6), Inches(3.7), Inches(4.8), "3. Object Storage (S3)", [
        "1 Bucket exclusivo na região Sudeste (br-se1).",
        "Visibilidade pública de leitura para entrega de mídias.",
        "Plugin Media Cloud Sync: salva imagens/vídeos direto no bucket.",
        "URLs das imagens servidas pelo endpoint S3 sem onerar a VM."
    ])
    add_footer(s3, 3)

    # =========================================================================
    # SLIDE 4: Por que desacoplar Mídias no Object Storage?
    # =========================================================================
    s4 = add_blank_slide()
    add_header(s4, "Benefícios da Arquitetura em Nuvem Desacoplada", "ARQUITETURA")
    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "❌ Arquitetura Monolítica Tradicional", [
        "Imagens e arquivos estáticos pesam no disco rígido da VM.",
        "Aumento contínuo de custo de armazenamento de bloco.",
        "Dificuldade de backup e recuperação de desastres.",
        "Impossibilidade de escalar a aplicação horizontalmente em múltiplas VMs."
    ])
    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "✅ Arquitetura Moderna na Magalu Cloud", [
        "VM permanece leve, usando disco mínimo apenas para SO e código.",
        "Object Storage com alta durabilidade, escalabilidade infinita e baixo custo.",
        "Mídias acessadas via URLs diretas do endpoint S3 da Magalu Cloud.",
        "Pronto para integração futura com CDN, cache global e clusters gerenciados."
    ])
    add_footer(s4, 4)

    # =========================================================================
    # SLIDE 5: Criação da Conta no ID Magalu
    # =========================================================================
    s5 = add_blank_slide()
    add_header(s5, "Passo 1: Criação da Conta no ID Magalu", "ONBOARDING")
    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "📝 Cadastro Individual", [
        "Cada participante utiliza sua própria conta no ID Magalu.",
        "Acesse: https://magalu.cloud e clique em 'Criar conta'.",
        "Preencha nome completo, CPF/CNPJ, e-mail e senha segura.",
        "Acesse sua caixa de entrada e confirme o e-mail de ativação.",
        "Documentação oficial: docs.magalu.cloud/docs/onboarding/create-account"
    ])
    add_card(s5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🔐 Primeiro Acesso ao Console", [
        "Acesse o Console: https://console.magalu.cloud",
        "Faça login com seu ID Magalu recém-criado.",
        "Aceite os Termos de Serviço da plataforma Magalu Cloud.",
        "Confirme a criação automática do seu Tenant (organização padrão).",
        "Pronto! Seu painel de controle está ativo para provisionamento."
    ])
    add_footer(s5, 5)

    # =========================================================================
    # SLIDE 6: Pré-requisitos Locais e Chaves SSH
    # =========================================================================
    s6 = add_blank_slide()
    add_header(s6, "Passo 2: Chave SSH e Ferramental Local", "PRÉ-REQUISITOS")
    add_card(s6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "🔑 Chave SSH Local", [
        "Necessária para autenticação e conexão remota na VM.",
        "Gere sua chave localmente via terminal:",
        "ssh-keygen -t ed25519 -f ~/.ssh/id_ed25519 -N '' -C 'workshop-mgc'",
        "Ajuste as permissões de segurança:",
        "chmod 600 ~/.ssh/id_ed25519",
        "Sua chave pública (~/.ssh/id_ed25519.pub) será cadastrada na MGC."
    ])
    add_card(s6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🛠️ Instalação das Ferramentas", [
        "O repositório inclui scripts automatizados de instalação:",
        "Linux / macOS:",
        "  chmod +x scripts/install-prereqs.sh && ./scripts/install-prereqs.sh",
        "Windows (PowerShell Administrador):",
        "  .\\scripts\\install-prereqs.ps1",
        "Instala automaticamente: mgc CLI, terraform, ansible e git."
    ])
    add_footer(s6, 6)

    # =========================================================================
    # SLIDE 7: Os 3 Métodos de Provisionamento
    # =========================================================================
    s7 = add_blank_slide()
    add_header(s7, "Escolha seu Método de Provisionamento", "OPÇÕES DE EXECUÇÃO")
    add_card(s7, Inches(0.8), Inches(1.6), Inches(3.6), Inches(4.8), "Método 1: Console Web", [
        "Interface gráfica intuitiva no navegador.",
        "Ideal para quem prefere uma experiência visual.",
        "Passo a passo por menus do Console da MGC.",
        "Deploy do WordPress via SSH manual.",
        "👉 Recomendado para iniciantes."
    ])
    add_card(s7, Inches(4.8), Inches(1.6), Inches(3.6), Inches(4.8), "Método 2: MGC CLI", [
        "Linha de comando oficial da Magalu Cloud.",
        "Execução rápida e direta no terminal.",
        "Comandos para VM, Chaves e Object Storage.",
        "Deploy do WordPress via SSH.",
        "👉 Recomendado para DevOps/SREs."
    ])
    add_card(s7, Inches(8.8), Inches(1.6), Inches(3.7), Inches(4.8), "Método 3: Terraform + Ansible", [
        "Infraestrutura como Código (IaC) 100% automatizada.",
        "Cria VM, Chave SSH, Bucket, ACL e IP público.",
        "Provisioner Ansible instala Docker e sobe o site.",
        "Deploy completo com 'terraform apply'.",
        "👉 Recomendado para automação avançada."
    ])
    add_footer(s7, 7)

    # =========================================================================
    # SLIDE 8: Procedimento 1 - Console (Chaves e Segurança)
    # =========================================================================
    s8 = add_blank_slide()
    add_header(s8, "Procedimento 1 (Console): Chaves SSH e Credenciais S3", "VIA CONSOLE WEB")
    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "1. Cadastrar Chave SSH no Console", [
        "No Console, acesse Perfil/Segurança → Chaves SSH.",
        "Clique em 'Adicionar chave SSH'.",
        "Nome da chave: site-key",
        "Cole o conteúdo da sua chave pública local (cat ~/.ssh/id_ed25519.pub).",
        "Clique em 'Salvar'.",
        "Doc: docs.magalu.cloud/docs/virtual-machine/quickstart"
    ])
    add_card(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "2. Gerar Chaves S3 do Object Storage", [
        "No menu lateral, acesse Object Storage → Chaves de Acesso.",
        "Clique em 'Criar chave de acesso' (escopo Leitura e Escrita).",
        "Copie e salve o Access Key ID e a Secret Access Key.",
        "Atenção: A Secret Key só é exibida uma vez no momento da criação!",
        "Doc: docs.magalu.cloud/docs/storage/object-storage/compatible-tools/mgc-cli-compatibility"
    ])
    add_footer(s8, 8)

    # =========================================================================
    # SLIDE 9: Procedimento 1 - Console (Criar VM)
    # =========================================================================
    s9 = add_blank_slide()
    add_header(s9, "Procedimento 1 (Console): Criar Máquina Virtual", "VIA CONSOLE WEB")
    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "📋 Configuração da Instância", [
        "Acesse: Máquinas Virtuais → Criar máquina.",
        "Nome: site-web-01",
        "Região: Sudeste (br-se1) | Zona: br-se1-a",
        "Imagem: Ubuntu 24.04 LTS (cloud-ubuntu-24.04 LTS)",
        "Configuração (Flavor): BV1-1-10 (1 vCPU / 1GB RAM) ou BV2-2-8",
        "Rede: VPC padrão (default)",
        "IP Público: Marcar opção para Alocar IP Público IPv4",
        "Chave SSH: Selecionar site-key"
    ])
    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🚀 Inicialização e Obtenção do IP", [
        "Clique em 'Criar Máquina Virtual'.",
        "Aguarde o status da máquina mudar de 'Criando' para 'Ativa'.",
        "Anote o IP Público alocado (ex: 201.23.xx.xx).",
        "Esse IP será utilizado para conectar via SSH e abrir o site no navegador.",
        "Doc: docs.magalu.cloud/docs/virtual-machine/quickstart"
    ])
    add_footer(s9, 9)

    # =========================================================================
    # SLIDE 10: Procedimento 1 - Console (Criar Bucket)
    # =========================================================================
    s10 = add_blank_slide()
    add_header(s10, "Procedimento 1 (Console): Criar Bucket no Object Storage", "VIA CONSOLE WEB")
    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "🪣 Configuração do Bucket", [
        "Acesse: Object Storage → Buckets → Criar bucket.",
        "Nome: site-media-<seu-nome> (deve ser único globalmente).",
        "Região: br-se1 (mesma região da VM).",
        "Visibilidade: Selecionar 'Público para leitura'.",
        "Clique em 'Criar bucket'.",
        "Doc: docs.magalu.cloud/docs/storage/object-storage/quickstart"
    ])
    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🌐 Acesso Público e Endpoint", [
        "O bucket está pronto para receber arquivos de mídia.",
        "Endpoint S3 da Magalu Cloud: https://br-se1.magaluobjects.com",
        "URL de acesso aos objetos do bucket:",
        "  https://br-se1.magaluobjects.com/site-media-<seu-nome>/<arquivo>",
        "Anote o nome do bucket para configurar o plugin no WordPress."
    ])
    add_footer(s10, 10)

    # =========================================================================
    # SLIDE 11: Procedimento 1 - Console (Deploy Docker na VM)
    # =========================================================================
    s11 = add_blank_slide()
    add_header(s11, "Procedimento 1 (Console): Deploy do WordPress via SSH", "VIA CONSOLE WEB")
    add_code_block(s11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "1. Conectar e Instalar Docker", [
        "# Conectar via SSH na VM",
        "ssh -i ~/.ssh/id_ed25519 ubuntu@<IP_PUBLICO>",
        "",
        "# Instalar Docker e Compose Plugin",
        "sudo apt update && sudo apt install -y curl ca-certificates",
        "sudo install -m 0755 -d /etc/apt/keyrings",
        "sudo curl -fsSL https://download.docker.com/linux/ubuntu/gpg \\",
        "  -o /etc/apt/keyrings/docker.asc",
        "echo \"deb [arch=$(dpkg --print-architecture) \\",
        "  signed-by=/etc/apt/keyrings/docker.asc] \\",
        "  https://download.docker.com/linux/ubuntu noble stable\" \\",
        "  | sudo tee /etc/apt/sources.list.d/docker.list",
        "sudo apt update && sudo apt install -y \\",
        "  docker-ce docker-ce-cli docker-compose-plugin"
    ])
    add_code_block(s11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "2. Subir WordPress com Docker Compose", [
        "mkdir -p ~/wordpress && cd ~/wordpress",
        "",
        "# Criar docker-compose.yml e iniciar",
        "cat > docker-compose.yml <<'EOF'",
        "services:",
        "  wordpress:",
        "    image: wordpress:latest",
        "    restart: always",
        "    ports: [\"80:80\"]",
        "    environment:",
        "      WORDPRESS_DB_HOST: db",
        "      WORDPRESS_DB_USER: wordpress",
        "      WORDPRESS_DB_PASSWORD: wp_password",
        "      WORDPRESS_DB_NAME: wordpress",
        "  db:",
        "    image: mysql:8.0",
        "    restart: always",
        "    environment:",
        "      MYSQL_DATABASE: wordpress",
        "      MYSQL_USER: wordpress",
        "      MYSQL_PASSWORD: wp_password",
        "      MYSQL_RANDOM_ROOT_PASSWORD: '1'",
        "EOF",
        "",
        "sudo docker compose up -d"
    ])
    add_footer(s11, 11)

    # =========================================================================
    # SLIDE 12: Procedimento 2 - CLI (Autenticação e Credenciais)
    # =========================================================================
    s12 = add_blank_slide()
    add_header(s12, "Procedimento 2 (CLI): Autenticação e Configurações", "VIA MGC CLI")
    add_code_block(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "1. Autenticar CLI e Configurar API Key", [
        "# Login via navegador no ID Magalu",
        "mgc auth login",
        "",
        "# Confirmar o tenant ativo",
        "mgc auth tenant current",
        "",
        "# Criar API Key para o Object Storage",
        "mgc object-storage api-key create workshop-key",
        "",
        "# Configurar a chave como padrão no CLI",
        "mgc object-storage api-key set <UUID_DA_KEY>",
        "mgc object-storage api-key current",
        "",
        "Doc: docs.magalu.cloud/docs/cli/"
    ])
    add_code_block(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "2. Registrar Chave SSH no Perfil", [
        "# Gerar chave local (se ainda não possuir)",
        "ssh-keygen -t ed25519 -f ~/.ssh/id_ed25519 \\",
        "  -N \"\" -C \"workshop-mgc\"",
        "",
        "# Registrar a chave pública no perfil da MGC",
        "mgc profile ssh-keys create \\",
        "  --name site-key \\",
        "  --key \"$(cat ~/.ssh/id_ed25519.pub)\" \\",
        "  --output json --raw",
        "",
        "# Listar chaves registradas",
        "mgc profile ssh-keys list",
        "",
        "Doc: docs.magalu.cloud/docs/cli/virtual-machine/instances"
    ])
    add_footer(s12, 12)

    # =========================================================================
    # SLIDE 13: Procedimento 2 - CLI (Criar VM e Obter IP)
    # =========================================================================
    s13 = add_blank_slide()
    add_header(s13, "Procedimento 2 (CLI): Criar Máquina Virtual", "VIA MGC CLI")
    add_code_block(s13, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8), "Comandos CLI para Criação e Consulta da VM", [
        "# 1. Criar a VM com IP público na região br-se1",
        "mgc virtual-machine instances create \\",
        "  --name site-web-01 \\",
        "  --machine-type.name BV1-1-10 \\",
        "  --image.name=\"cloud-ubuntu-24.04 LTS\" \\",
        "  --ssh-key-name site-key \\",
        "  --network.associate-public-ip true \\",
        "  --region br-se1 \\",
        "  --output json --raw",
        "",
        "# 2. Listar instâncias e copiar o IP Público associado",
        "mgc virtual-machine instances list",
        "",
        "# 3. Consultar detalhes da instância criada",
        "mgc virtual-machine instances get <ID_DA_INSTANCIA>",
        "",
        "Doc oficial: docs.magalu.cloud/docs/cli/virtual-machine/instances"
    ])
    add_footer(s13, 13)

    # =========================================================================
    # SLIDE 14: Procedimento 2 - CLI (Criar Bucket e Política)
    # =========================================================================
    s14 = add_blank_slide()
    add_header(s14, "Procedimento 2 (CLI): Criar Bucket no Object Storage", "VIA MGC CLI")
    add_code_block(s14, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8), "Comandos CLI para Criação e Configuração de Leitura Pública do Bucket", [
        "# 1. Criar o bucket na região Sudeste (br-se1)",
        "mgc object-storage buckets create site-media-<seu-nome> \\",
        "  --public-read \\",
        "  --region br-se1 \\",
        "  --output json --raw",
        "",
        "# 2. Definir a política de leitura pública (s3:GetObject) para entrega de mídias",
        "mgc object-storage buckets policy set \\",
        "  --dst site-media-<seu-nome> \\",
        "  --policy '{",
        "    \"Version\": \"2012-10-17\",",
        "    \"Statement\": [{",
        "      \"Effect\": \"Allow\",",
        "      \"Principal\": \"*\",",
        "      \"Action\": \"s3:GetObject\",",
        "      \"Resource\": \"site-media-<seu-nome>/*\"",
        "    }]",
        "  }'",
        "",
        "Doc oficial: docs.magalu.cloud/docs/storage/object-storage/compatible-tools/mgc-cli-compatibility"
    ])
    add_footer(s14, 14)

    # =========================================================================
    # SLIDE 15: Procedimento 3 - Terraform (Visão Geral IaC)
    # =========================================================================
    s15 = add_blank_slide()
    add_header(s15, "Procedimento 3 (IaC): Automação com Terraform e Ansible", "TERRAFORM + ANSIBLE")
    add_card(s15, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "🏗️ O que o Terraform Provisiona?", [
        "Chave SSH (mgc_ssh_keys): registra a chave pública na nuvem.",
        "VM (mgc_virtual_machine_instances): aloca a VM com IP público.",
        "Bucket (mgc_object_storage_buckets): cria o repositório de mídias.",
        "ACL S3 (aws_s3_bucket_acl): aplica permissão public-read via endpoint S3.",
        "Outputs: exibe IP público, comando SSH e URL do site no terminal.",
        "Doc: docs.magalu.cloud/docs/infrastructure-as-code/terraform"
    ])
    add_card(s15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🤖 O que o Ansible Executa Automaticamente?", [
        "Trigger local-exec é acionado logo após a VM ficar disponível.",
        "Gera o inventário hosts.ini dinamicamente com o IP da VM.",
        "Aguarda o daemon SSH e o boot da VM (cloud-init status --wait).",
        "Instala o Docker Engine e o plugin Docker Compose.",
        "Copia o docker-compose.yml e inicia os containers do WordPress.",
        "Zero intervenção manual necessária!"
    ])
    add_footer(s15, 15)

    # =========================================================================
    # SLIDE 16: Procedimento 3 - Terraform (Execução Passo a Passo)
    # =========================================================================
    s16 = add_blank_slide()
    add_header(s16, "Procedimento 3 (IaC): Executando o Terraform", "TERRAFORM + ANSIBLE")
    add_code_block(s16, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "1. Configurar terraform.tfvars", [
        "cd terraform",
        "cp terraform.tfvars.example terraform.tfvars",
        "",
        "# Preencher variáveis obrigatórias:",
        "mgc_api_key          = \"SUA_API_KEY_MGC\"",
        "region               = \"br-se1\"",
        "vm_name              = \"site-web-01\"",
        "ssh_key_name         = \"site-key\"",
        "ssh_public_key       = \"ssh-ed25519 AAAAC3...\"",
        "ssh_private_key_path = \"~/.ssh/id_ed25519\"",
        "bucket_name          = \"site-media-<seu-nome>\"",
        "object_storage_access_key = \"SUA_ACCESS_KEY\"",
        "object_storage_secret_key = \"SUA_SECRET_KEY\""
    ])
    add_code_block(s16, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "2. Aplicar o Provisionamento", [
        "# Inicializar providers (MGC, AWS S3, Null)",
        "terraform init",
        "",
        "# Validar o plano de execução",
        "terraform plan",
        "",
        "# Aplicar a infraestrutura e rodar o deploy",
        "terraform apply -auto-approve",
        "",
        "# Ao final, veja os outputs no terminal:",
        "Outputs:",
        "  bucket_name = \"site-media-seu-nome\"",
        "  public_ip   = \"201.23.xx.xx\"",
        "  wp_url      = \"http://201.23.xx.xx\""
    ])
    add_footer(s16, 16)

    # =========================================================================
    # SLIDE 17: Configuração Inicial do WordPress
    # =========================================================================
    s17 = add_blank_slide()
    add_header(s17, "Finalizando a Instalação do WordPress", "CONFIGURAÇÃO DO SITE")
    add_card(s17, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "🌐 1. Assistente Web Inicial", [
        "Abra o navegador no endereço: http://<IP_PUBLICO>",
        "Selecione o idioma: Português do Brasil.",
        "Preencha as informações do site:",
        "  • Título do Site (ex: 'Meu Site na Magalu Cloud')",
        "  • Nome de Usuário Administrador (ex: admin)",
        "  • Senha de acesso e E-mail",
        "Clique em 'Instalar WordPress'."
    ])
    add_card(s17, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🔐 2. Login no Painel Administrativo", [
        "Após concluir a instalação, acesse a tela de login:",
        "  http://<IP_PUBLICO>/wp-admin",
        "Informe seu usuário e senha recém-criados.",
        "Você terá acesso ao painel de controle completo do WordPress.",
        "Próximo passo: conectar a biblioteca de mídias ao Object Storage da Magalu Cloud."
    ])
    add_footer(s17, 17)

    # =========================================================================
    # SLIDE 18: Conectar Mídias ao Object Storage (Plugin)
    # =========================================================================
    s18 = add_blank_slide()
    add_header(s18, "Integrando o WordPress ao Object Storage S3", "CONFIGURAÇÃO DO SITE")
    add_card(s18, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "🔌 1. Instalar Plugin Media Cloud Sync", [
        "No painel do WordPress, vá em Plugins → Adicionar Novo.",
        "Busque por: Media Cloud Sync",
        "Clique em 'Instalar Agora' e depois em 'Ativar'.",
        "Acesse as configurações do plugin (menu lateral Media Cloud ou Settings → Cloud Storage)."
    ])
    add_card(s18, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "⚙️ 2. Configurar Credenciais S3 da MGC", [
        "Selecione o provedor: S3 Compatible",
        "Provider Label: Magalu Cloud",
        "Storage Endpoint: https://br-se1.magaluobjects.com/",
        "Access Key & Secret Key: suas chaves do Object Storage",
        "Region: br-se1",
        "Bucket Name: site-media-<seu-nome>",
        "Media Delivery Provider: Other / Magalu Cloud",
        "Clique em 'Save Changes' e teste a conexão."
    ])
    add_footer(s18, 18)

    # =========================================================================
    # SLIDE 19: Validação e Teste Hands-on
    # =========================================================================
    s19 = add_blank_slide()
    add_header(s19, "Validação Hands-on do Ambiente", "TESTE & VERIFICAÇÃO")
    add_card(s19, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "📸 1. Upload de Imagem no WordPress", [
        "No painel administrativo, acesse Mídia → Adicionar Nova.",
        "Faça o upload de uma imagem ou foto de teste (.jpg / .png).",
        "Abra os detalhes da mídia e copie o link gerado ('URL do arquivo').",
        "Observe que a URL aponta diretamente para:",
        "https://br-se1.magaluobjects.com/site-media-<seu-nome>/wp-content/uploads/..."
    ])
    add_card(s19, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🔍 2. Prova de Acesso e Entrega S3", [
        "Abra uma aba anônima no navegador e cole a URL da imagem.",
        "A imagem abre perfeitamente, servida direto do Object Storage da MGC!",
        "Confira também no Console da Magalu Cloud que o arquivo está no bucket.",
        "A VM não guarda o arquivo estático no disco: economia de storage e alta durabilidade!"
    ])
    add_footer(s19, 19)

    # =========================================================================
    # SLIDE 20: Desmontagem e Exclusão dos Recursos
    # =========================================================================
    s20 = add_blank_slide()
    add_header(s20, "Desmontagem dos Recursos (Sem Custo Residual)", "ENCERRAMENTO")
    add_card(s20, Inches(0.8), Inches(1.6), Inches(3.6), Inches(4.8), "Exclusão no Console", [
        "1. Máquinas Virtuais:",
        "  Selecionar site-web-01 → Excluir (libera o IP público).",
        "2. Object Storage:",
        "  Selecionar o bucket → Esvaziar objetos → Excluir bucket.",
        "3. Chaves SSH:",
        "  Segurança → Excluir site-key."
    ])
    add_card(s20, Inches(4.8), Inches(1.6), Inches(3.6), Inches(4.8), "Exclusão via MGC CLI", [
        "# Obter ID e deletar VM + IP:",
        "VM_ID=$(mgc vm instances list -o json | grep -o '\"id\": \"[^\"]*' | head -1 | cut -d'\"' -f4)",
        "mgc vm instances delete --id $VM_ID --delete-public-ip=true",
        "# Deletar bucket e arquivos:",
        "mgc object-storage buckets delete --dst site-media-<nome> --recursive",
        "# Deletar chave SSH:",
        "mgc profile ssh-keys delete --name site-key"
    ])
    add_card(s20, Inches(8.8), Inches(1.6), Inches(3.7), Inches(4.8), "Exclusão via Terraform", [
        "# No diretório terraform/:",
        "cd terraform",
        "terraform destroy -auto-approve",
        "",
        "O Terraform remove automaticamente:",
        "  • Máquina Virtual e IP Público",
        "  • Bucket e ACL do Object Storage",
        "  • Chave SSH registrada",
        "Ambiente 100% limpo!"
    ])
    add_footer(s20, 20)

    # =========================================================================
    # SLIDE 21: Principais Aprendizados e Próximos Passos
    # =========================================================================
    s21 = add_blank_slide()
    add_header(s21, "Principais Aprendizados e Evoluções Arquiteturais", "SUMÁRIO")
    add_card(s21, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.8), "🎓 O que Aprendemos Hoje?", [
        "Autenticação e gestão de contas com o ID Magalu.",
        "Provisionamento de instâncias de VM na nuvem Magalu.",
        "Uso do Object Storage S3 para armazenamento escalável de mídias.",
        "Gerenciamento de recursos via Console, CLI e Infraestrutura como Código (Terraform).",
        "Deploy automatizado de aplicações web com Docker Compose e Ansible."
    ])
    add_card(s21, Inches(6.8), Inches(1.6), Inches(5.7), Inches(4.8), "🚀 Como Evoluir Essa Arquitetura?", [
        "DBaaS: Migrar o banco MySQL para o Database-as-a-Service gerenciado da MGC.",
        "Kubernetes (MKE): Orquestrar a aplicação em containers escaláveis.",
        "Load Balancer: Distribuir tráfego entre múltiplas instâncias web.",
        "CDN & WAF: Acelerar entrega de conteúdo e proteger contra ameaças.",
        "Backups e Políticas de Ciclo de Vida no Object Storage."
    ])
    add_footer(s21, 21)

    # =========================================================================
    # SLIDE 22: Encerramento e Links Úteis
    # =========================================================================
    s22 = add_blank_slide()
    add_header(s22, "Obrigado e Próximos Passos!", "DOCUMENTAÇÃO & CONTATOS")
    add_card(s22, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8), "📚 Links Úteis e Documentações Oficiais da Magalu Cloud", [
        "🌐 Portal Oficial: https://magalu.cloud/",
        "📖 Documentação Central: https://docs.magalu.cloud/",
        "🚀 Guia de Criação de Conta: https://docs.magalu.cloud/docs/onboarding/create-account",
        "💻 Guia da MGC CLI: https://docs.magalu.cloud/docs/cli/installation-and-configuration/get-started",
        "🖥️ Máquinas Virtuais (Compute): https://docs.magalu.cloud/docs/virtual-machine/quickstart",
        "🪣 Object Storage S3: https://docs.magalu.cloud/docs/storage/object-storage/quickstart",
        "🏗️ Terraform Provider MGC: https://registry.terraform.io/providers/MagaluCloud/mgc/latest/docs",
        "🐙 Repositório do Workshop: https://github.com/fernandoravagnani-luizalabs/mgc-workshop"
    ])
    add_footer(s22, 22)

    output_path = "/Users/fp_ravagnani/Documents/Workshop/mgc-workshop/Workshop-Magalu-Cloud.pptx"
    prs.save(output_path)
    print(f"Apresentação gerada com sucesso em: {output_path}")

if __name__ == "__main__":
    create_presentation()
